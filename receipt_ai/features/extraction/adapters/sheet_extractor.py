from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from receipt_ai.features.extraction.config import ExtractionConfig
from receipt_ai.features.extraction.cancel_checks import raise_if_cancelled
from receipt_ai.features.extraction.contracts import Extractor
from receipt_ai.features.extraction.errors import ExtractionCancelledError, ParsingError
from receipt_ai.features.extraction.models import ExtractionRequest


class SheetExtractor(Extractor):
    name = "csv-openpyxl"

    def __init__(self, config: ExtractionConfig) -> None:
        self._config = config

    def supports(self, request: ExtractionRequest) -> bool:
        lowered = request.filename.lower()
        return lowered.endswith(".csv") or lowered.endswith(".xlsx")

    def extract(self, request: ExtractionRequest) -> str:
        lowered = request.filename.lower()
        if lowered.endswith(".csv"):
            return self._extract_csv(request)
        return self._extract_xlsx(request)

    def _extract_csv(self, request: ExtractionRequest) -> str:
        try:
            text = request.file_bytes.decode("utf-8", errors="replace")
            rows = list(csv.reader(io.StringIO(text)))
            lines: list[str] = []
            for ri, row in enumerate(rows[: self._config.max_sheet_rows]):
                if ri % 50 == 0:
                    raise_if_cancelled(request)
                lines.append(" | ".join(cell.strip() for cell in row[: self._config.max_sheet_cols]))
            return "\n".join(lines).strip()
        except ExtractionCancelledError:
            raise
        except Exception as exc:
            raise ParsingError(f"CSV parse failed: {exc}") from exc

    def _extract_xlsx(self, request: ExtractionRequest) -> str:
        try:
            workbook = load_workbook(io.BytesIO(request.file_bytes), read_only=True, data_only=True)
            parts: list[str] = []
            try:
                for sheet in workbook.worksheets:
                    parts.append(f"## {sheet.title}")
                    row_idx = 0
                    for row in sheet.iter_rows(values_only=True):
                        if row_idx % 50 == 0:
                            raise_if_cancelled(request)
                        if row_idx >= self._config.max_sheet_rows:
                            parts.append(f"... truncated at {self._config.max_sheet_rows} rows")
                            break
                        cells = ["" if cell is None else str(cell) for cell in row[: self._config.max_sheet_cols]]
                        if any(cell.strip() for cell in cells):
                            parts.append(" | ".join(cells))
                        row_idx += 1
                    parts.append("")
            finally:
                workbook.close()
            return "\n".join(parts).strip()
        except ExtractionCancelledError:
            raise
        except Exception as exc:
            raise ParsingError(f"XLSX parse failed: {exc}") from exc

from __future__ import annotations

import io

import docx

from receipt_ai.features.extraction.cancel_checks import raise_if_cancelled
from receipt_ai.features.extraction.contracts import Extractor
from receipt_ai.features.extraction.errors import ExtractionCancelledError, ParsingError
from receipt_ai.features.extraction.models import ExtractionRequest


class DocxExtractor(Extractor):
    name = "python-docx"

    def supports(self, request: ExtractionRequest) -> bool:
        return request.filename.lower().endswith(".docx")

    def extract(self, request: ExtractionRequest) -> str:
        try:
            parsed = docx.Document(io.BytesIO(request.file_bytes))
            lines: list[str] = []
            for pi, paragraph in enumerate(parsed.paragraphs):
                if pi % 25 == 0:
                    raise_if_cancelled(request)
                text = paragraph.text.strip()
                if text:
                    lines.append(text)
            for ti, table in enumerate(parsed.tables):
                if ti % 5 == 0:
                    raise_if_cancelled(request)
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    if any(row_text):
                        lines.append(" | ".join(row_text))
            return "\n".join(lines).strip()
        except ExtractionCancelledError:
            raise
        except Exception as exc:
            raise ParsingError(f"DOCX parse failed: {exc}") from exc
